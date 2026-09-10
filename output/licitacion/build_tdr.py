from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path(__file__).parent
d=Document(); s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
s.header_distance=s.footer_distance=Inches(.492)
# Preset standard_business_brief; memo_masthead. Named overrides: title 26pt;
# compact table text 9pt; page-level section breaks for reviewable specification.
for name,size,color,before,after in [('Normal',11,'222222',0,6),('Title',26,'0B2545',0,10),('Subtitle',12,'555555',0,8),('Heading 1',16,'2E74B5',16,8),('Heading 2',13,'2E74B5',12,6),('Heading 3',12,'1F4D78',8,4)]:
 st=d.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color)
 st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.line_spacing=1.10
 st.paragraph_format.widow_control=True
 if name.startswith('Heading'): st.paragraph_format.keep_with_next=True
for name in ['Header','Footer']:
 d.styles[name].font.name='Calibri'; d.styles[name].font.size=Pt(9); d.styles[name].font.color.rgb=RGBColor.from_string('666666')
s.header.paragraphs[0].text='SISTEMA INTEGRAL DE SERVICIOS DE SALUD  |  TDR · BORRADOR'
f=s.footer.paragraphs[0]; f.alignment=2
f.add_run('Versión 1.0 · 07 septiembre 2026  |  ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)
def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,2)
def page(t): d.add_page_break(); d.add_heading(t,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers)); t.autofit=False
 pr=t._tbl.tblPr
 for tag,val in [('tblW',9360),('tblInd',120)]:
  el=pr.find(qn('w:'+tag))
  if el is None: el=OxmlElement('w:'+tag); pr.append(el)
  el.set(qn('w:w'),str(val)); el.set(qn('w:type'),'dxa')
 mar=OxmlElement('w:tblCellMar')
 for k,v in [('top',80),('bottom',80),('start',120),('end',120)]:
  el=OxmlElement('w:'+k); el.set(qn('w:w'),str(v)); el.set(qn('w:type'),'dxa'); mar.append(el)
 pr.append(mar)
 borders=OxmlElement('w:tblBorders')
 for k in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+k); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'CCCCCC'); borders.append(el)
 pr.append(borders)
 for c,w in zip(t._tbl.tblGrid.gridCol_lst,widths): c.set(qn('w:w'),str(w))
 for row in rows: t.add_row()
 for i,(row,vals) in enumerate(zip(t.rows,[headers]+rows)):
  trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
  if i==0: trpr.append(OxmlElement('w:tblHeader'))
  for c,txt,w in zip(row.cells,vals,widths):
   c.width=Inches(w/1440); c._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),str(w)); c.text=txt
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(3); para.paragraph_format.space_before=Pt(0)
    for r in para.runs: r.font.size=Pt(9); r.bold=i==0
   if i==0:
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'F2F4F7'); c._tc.get_or_add_tcPr().append(sh)

d.add_paragraph('TÉRMINOS DE REFERENCIA',style='Title')
d.add_paragraph('Contratación de una plataforma integral de gestión clínica, hospitalaria y de salud pública',style='Subtitle')
p('Suministro, parametrización, implementación, migración, capacitación y soporte técnico de una solución para una red de establecimientos de salud.')
h('Datos del proceso')
for t in ['Entidad contratante: [COMPLETAR]','Unidad solicitante: [COMPLETAR]','Número y modalidad del proceso: [COMPLETAR]','Fuente de financiamiento y presupuesto: [COMPLETAR]','Cobertura territorial y establecimientos: [COMPLETAR]','Plazo propuesto: 24 semanas de implementación y 12 meses de garantía y soporte desde la recepción final.'] : p(t)
h('Estado y uso del documento')
p('Borrador técnico para estructuración de la contratación. Los plazos, niveles de servicio y ponderaciones son propuestas para validación institucional, no condiciones de un proceso ya aprobado. Los campos [COMPLETAR] deben resolverse antes de publicar.')
p('El alcance se basa en la documentación y estructura del Sistema Integral de Servicios de Salud (SISS). Las especificaciones se expresan por resultados y admiten soluciones equivalentes. La presencia de módulos en el repositorio no constituye certificación de funcionamiento, seguridad ni despliegue productivo.')
p('El anexo final es de uso interno para preparar la oferta del SISS y debe retirarse de la versión que se someta como términos de referencia.')

page('1. Antecedentes, objeto y cobertura')
p('La entidad requiere integrar la identificación del paciente, programación de servicios, atención clínica, apoyo diagnóstico, suministro de medicamentos y seguimiento de salud pública. Se busca disminuir la recaptura, mejorar la continuidad asistencial y disponer de información trazable por establecimiento y a nivel consolidado.')
h('Objeto de la contratación')
p('Contratar una solución informática integral y los servicios necesarios para ponerla en operación, incluyendo configuración de procesos, interfaces acordadas, carga de información, pruebas, capacitación, documentación, garantía y acompañamiento operativo.')
h('Resultados esperados')
p('Expediente longitudinal vinculado a la atención; agendas coordinadas entre sedes; trazabilidad de recetas y existencias por lote; registros de vacunación y vigilancia; indicadores consistentes y acceso ciudadano limitado a la información autorizada.')
h('Dimensionamiento previo a publicación')
table(['Parámetro','Dato institucional requerido'],[
['Establecimientos y fases','[COMPLETAR] cantidad, ubicación, servicios y sedes piloto.'],
['Usuarios y demanda','[COMPLETAR] usuarios nominales, concurrentes, consultas y recetas por día.'],
['Migración','[COMPLETAR] fuentes, años, registros, adjuntos, calidad y volumen en GB.'],
['Infraestructura','[COMPLETAR] alojamiento, conectividad, equipos y responsabilidades.'],
['Integraciones','[COMPLETAR] sistemas externos, interfaces, responsables y autorizaciones.'],
['Servicio','[COMPLETAR] horario operativo, criticidad y sedes con atención continua.']], [2400,6960])
h('Límites de alcance')
p('Se incluyen los requisitos funcionales de la sección 2 y los servicios de implementación. Equipos médicos, hardware, conectividad, mensajería de pago, licencias de terceros y nuevas integraciones deberán identificarse y cotizarse por separado cuando sean necesarios. PACS, equipos de laboratorio y consultas a registros externos requieren definición expresa de interfaz; no se consideran incluidos por una mención genérica.')

page('2. Requisitos funcionales · atención')
p('Todos los requisitos RF son obligatorios en la recepción final. El oferente declarará para cada uno si está disponible, requiere parametrización, requiere desarrollo o depende de un tercero, indicando evidencia y fecha de entrega.')
for title,txt in [
('RF-01. Red de establecimientos y administración','Gestionar establecimientos, servicios, especialidades, personal, roles y permisos por ámbito. Permitir que un profesional atienda en varias sedes sin habilitar acceso indiscriminado a datos. Validación: un mismo usuario opera en dos sedes autorizadas y se rechaza su acceso a una tercera.'),
('RF-02. Identificación y expediente del paciente','Registrar datos demográficos, identificadores, contactos y antecedentes; detectar potenciales duplicados y conservar trazabilidad de las correcciones. Consultar el historial longitudinal según permisos. Validación: búsqueda por identificador, advertencia de duplicado y consulta del historial autorizado.'),
('RF-03. Agendas y citas','Configurar jornadas, duración de citas, cupos y excepciones; prevenir conflictos del profesional entre establecimientos. Registrar programación, reprogramación, cancelación y estados de atención. Validación: intentar una doble reserva y una reserva durante una ausencia, ambas rechazadas.'),
('RF-04. Triaje y consulta clínica','Registrar signos vitales, antecedentes y notas SOAP; vincular diagnósticos CIE-10, prescripciones, órdenes, referencias e incapacidades según configuración institucional. Validación: completar una consulta y recuperar todos sus documentos vinculados al paciente y al profesional.'),
('RF-05. Formularios clínicos configurables','Permitir formularios por servicio o especialidad, con campos requeridos y conservación de la versión utilizada. Validación: modificar una plantilla sin alterar los registros históricos capturados con la versión anterior.'),
('RF-06. Hospitalización y continuidad','Gestionar disponibilidad y asignación de camas, ingresos, movimientos y egresos; vincular evolución y referencias al expediente. Validación: ingreso, traslado y egreso de un paciente sin doble ocupación de una cama.'),
('RF-07. Atención materna e infantil','Registrar controles prenatales y seguimiento pediátrico, con variables, antecedentes y evolución definidos por la institución. Validación: dos controles sucesivos permiten consultar la evolución y emitir el reporte correspondiente.')]:
 h(title); p(txt)

page('2. Requisitos funcionales · suministros y salud pública')
for title,txt in [
('RF-08. Prescripción y dispensación','Registrar dosis, vía, frecuencia, duración y cantidad; gestionar entregas completas y parciales, cancelaciones y demanda insatisfecha. Identificar tratamientos coincidentes para revisión profesional. Validación: dispensar parcialmente y conservar el saldo, sin duplicar el descargo al repetir la solicitud.'),
('RF-09. Inventario farmacéutico','Controlar entradas, salidas, ajustes y bajas por establecimiento, lote y vencimiento. Aplicar FEFO, impedir existencias negativas y uso de lotes vencidos; registrar usuario y motivo. Validación: dos dispensaciones simultáneas sobre un saldo limitado no generan sobreventa ni inconsistencias.'),
('RF-10. Programa de inmunizaciones','Configurar biológicos y esquemas aprobados por la institución; registrar aplicación, dosis, lote, fecha y vacunador; controlar existencias y pérdidas, incluidas las atribuidas a cadena de frío; generar carnet. Validación: rastrear una aplicación hasta su lote y conciliar el inventario. La telemetría de temperatura no está implícita.'),
('RF-11. Vigilancia epidemiológica','Relacionar diagnósticos configurados con alertas y fichas de seguimiento, semana epidemiológica, ubicación y estados de gestión. Restringir el acceso a datos identificables. Validación: un diagnóstico configurado genera el registro previsto y aparece en el consolidado autorizado.'),
('RF-12. Laboratorio e imagenología','Vincular solicitudes, estados, resultados, unidades y rangos de referencia cuando correspondan; registrar informes radiológicos y adjuntos o enlaces autorizados. Validación: el resultado validado se consulta desde el encuentro clínico. Las conexiones con analizadores o PACS se aceptan mediante pruebas de interfaz específicas.'),
('RF-13. Servicios al paciente','Disponer de un canal móvil o web adaptable para citas, recetas, carnet y consultas personales autorizadas. Si se publica disponibilidad de medicamentos, mostrar sede y fecha de actualización sin comprometer reservas. Validación: un paciente no puede consultar registros de otro mediante la interfaz ni la API.'),
('RF-14. Reportes y trazabilidad','Generar reportes por fecha, establecimiento y servicio sobre atenciones, inventario, demanda insatisfecha, vacunación y vigilancia; exportar datos autorizados y documentos PDF. Validación: los totales coinciden con un conjunto de datos de control y las operaciones críticas pueden rastrearse.')]:
 h(title); p(txt)

page('2.1 Componente móvil para pacientes')
p('La solución incluirá una aplicación móvil para pacientes, integrada con la misma información clínica y administrativa autorizada de la plataforma central. El componente deberá funcionar en Android e iOS, mantener una experiencia coherente en teléfonos de gama media y aplicar controles específicos para datos personales y de salud. El oferente indicará versiones mínimas de sistema operativo, dispositivos probados y requisitos de conectividad.')
for title,txt in [
('RM-01. Registro y acceso del paciente','Permitir solicitud de cuenta mediante datos de identificación y contacto definidos por la institución, activación conforme al flujo aprobado, inicio y cierre de sesión, recuperación o cambio de contraseña y cambio obligatorio de credenciales temporales. Validación: una solicitud queda pendiente de aprobación, una cuenta autorizada accede únicamente a su perfil y una sesión cerrada no reutiliza el token anterior.'),
('RM-02. Inicio y perfil personal','Presentar un resumen personal con próxima cita, accesos a los servicios habilitados y datos básicos del expediente. Permitir actualización solo de los campos autorizados y mostrar el origen o fecha de actualización cuando ayude a evitar decisiones con datos desactualizados. Validación: el paciente visualiza su resumen y no puede alterar datos clínicos protegidos.'),
('RM-03. Autogestión de citas','Consultar establecimientos, especialidades, profesionales y horarios disponibles; solicitar una cita y recibir confirmación con fecha, hora, sede, especialidad, profesional y estado. Aplicar cupos, restricciones y prevención de duplicidades definidas por la institución. Validación: crear una cita disponible, rechazar un horario ocupado y reflejar inmediatamente la reserva en la agenda institucional.'),
('RM-04. Historial clínico personal','Consultar episodios autorizados del historial, incluyendo fecha, establecimiento, profesional, diagnósticos y resumen permitido por la política institucional. La app no deberá exponer notas o datos restringidos. Validación: el contenido coincide con el expediente central y las pruebas de acceso horizontal impiden consultar a otro paciente.'),
('RM-05. Recetas y medicamentos','Consultar recetas vigentes e históricas, estado, establecimiento, medicamentos, dosis, frecuencia, duración y cantidad autorizada. Validación: una dispensación parcial o completa actualiza el estado visible sin modificar la prescripción original.'),
('RM-06. Disponibilidad por establecimiento','Buscar medicamentos por nombre o código y mostrar disponibilidad por establecimiento con fecha y hora de actualización. La interfaz deberá advertir que la disponibilidad es informativa cuando no exista reserva. Validación: la búsqueda coincide con saldos publicables de un conjunto de control y no revela lotes, costos ni información operativa restringida.'),
('RM-07. Carnet digital y vacunación','Mostrar identificación del paciente y carnet digital con un código verificable de vigencia limitada o mecanismo equivalente que evite exponer directamente un identificador sensible. Consultar historial de vacunas autorizado y generar o descargar el carnet institucional cuando corresponda. Validación: un código vencido o alterado no permite obtener información y el carnet coincide con las dosis registradas.'),
('RM-08. Seguimiento prenatal','Cuando aplique, mostrar embarazo activo, fecha probable de parto, clasificación de riesgo, controles, signos y observaciones autorizadas; permitir descargar el documento institucional. Validación: la información coincide con el módulo clínico y un embarazo cerrado deja de presentarse como activo.'),
('RM-09. Avisos y notificaciones','Enviar recordatorios de citas y avisos transaccionales aprobados mediante notificaciones push o mecanismo equivalente, con consentimiento y configuración del usuario. El texto visible en pantalla bloqueada no deberá revelar diagnósticos, medicamentos ni otros datos sensibles. Validación: registrar entrega o fallo de una notificación de prueba y respetar la desactivación del usuario.'),
('RM-10. Experiencia, accesibilidad y conectividad','Ofrecer interfaz en español, navegación comprensible, mensajes de error útiles, tamaños de texto adaptables y compatibilidad con lectores de pantalla en los flujos principales. Ante pérdida de red, conservar de forma segura solo la información necesaria, impedir envíos duplicados y reintentar de manera controlada. El alcance de operación sin conexión se definirá antes de publicar. Validación: completar los casos principales con el tamaño de fuente ampliado y recuperar una operación interrumpida sin duplicarla.'),
('RM-11. Distribución y administración de versiones','Entregar paquetes instalables firmados, configuración por ambiente, identificadores institucionales y procedimiento de actualización. Publicar o distribuir por los canales aprobados por la entidad y transferir cuentas, certificados, llaves y registros de publicación según el régimen contractual. Validación: instalar una versión de entrega en dispositivos autorizados y actualizarla preservando la sesión o solicitando reautenticación de forma controlada.'),
('RM-12. Analítica y soporte móvil','Registrar fallos técnicos y métricas operativas mínimas sin capturar datos clínicos en servicios de analítica. Mostrar versión de la app, canal de soporte y condiciones de privacidad. Validación: un incidente aporta versión y diagnóstico técnico suficiente sin incluir credenciales, tokens ni contenido clínico.')]:
 h(title); p(txt)

page('3. Requisitos técnicos y protección de información')
for title,txt in [
('RT-01. Arquitectura y mantenimiento','Solución modular, interfaz en español, configuración de zona horaria institucional y API documentada. Se aceptan tecnologías equivalentes; no se exige un fabricante, lenguaje o marco específico. Entregar arquitectura, dependencias, versiones soportadas y procedimiento de actualización.'),
('RT-02. Acceso y auditoría','Aplicar autenticación segura, permisos en servidor, mínimo privilegio, cierre de sesión y revocación de accesos; autenticación multifactor para cuentas privilegiadas. Registrar accesos y cambios críticos con usuario, fecha, sede y acción; proteger las bitácoras frente a alteraciones no autorizadas.'),
('RT-03. Datos sensibles','Cifrar comunicaciones y respaldos; proteger credenciales y secretos; separar ambientes de desarrollo, pruebas y producción. Usar datos sintéticos o anonimizados en demostraciones. Definir conservación, acceso, exportación y eliminación con la entidad antes de la operación.'),
('RT-04. Integridad y continuidad','Garantizar consistencia transaccional de inventario y atención, copias de seguridad automatizadas y restauración verificable. Objetivos propuestos: RPO máximo 24 horas y RTO máximo 8 horas, sujetos a dimensionamiento. Demostrar restauración en ambiente aislado con conciliación de registros.'),
('RT-05. Rendimiento y disponibilidad','Objetivo propuesto: percentil 95 de hasta 3 segundos en consultas y registros ordinarios, excluyendo reportes masivos e interfaces externas, con [COMPLETAR] usuarios concurrentes y volumen acordado. Ejecutar una prueba de 60 minutos tras calentamiento y entregar métricas de errores y tiempos. Disponibilidad mensual propuesta: 99.5% en la ventana contratada.'),
('RT-06. Interoperabilidad y salida','Documentar contratos de API, autenticación, errores, catálogos y formatos de intercambio. Para cada integración, fijar origen, destino, campos, frecuencia y prueba de aceptación. Entregar exportación completa de datos y adjuntos en formatos abiertos, con diccionario y relaciones suficientes para reutilización.'),
('RT-07. Verificación técnica','Entregar pruebas de autorización, integridad, restauración y rendimiento; análisis de vulnerabilidades y correcciones. La recepción productiva exige ausencia de hallazgos críticos o altos pendientes que afecten el alcance, clasificados con un método acordado. Los resultados de pruebas deben indicar versión, ambiente y fecha.'),
('RT-08. Seguridad del componente móvil','Usar exclusivamente HTTPS en producción, validar certificados y deshabilitar tráfico en texto claro. Almacenar tokens, llaves y secretos en mecanismos seguros del sistema operativo; no incluir secretos, credenciales ni URL locales fijas en el paquete. Aplicar expiración y revocación de sesiones, ocultar datos sensibles en registros y vistas recientes, minimizar permisos y proteger archivos descargados. Someter las versiones Android e iOS a pruebas de seguridad móvil basadas en un estándar reconocido y corregir hallazgos críticos o altos antes de la recepción.'),
('RT-09. Calidad y compatibilidad móvil','Probar instalación, actualización, accesibilidad, consumo de red, recuperación ante interrupciones y flujos RM en una matriz de dispositivos reales y emulados aprobada. La entidad completará las versiones mínimas de Android [COMPLETAR] e iOS [COMPLETAR]. Las compilaciones de producción deberán usar identificador, firma y configuración institucional, sin claves de depuración.'),
('RT-10. Servicios de publicación y mensajería','El oferente configurará los servicios de notificaciones y publicación definidos en el alcance, documentará sus costos y dependencias y transferirá la administración a la entidad. Los ambientes de desarrollo, pruebas y producción utilizarán credenciales separadas. La aplicación deberá seguir operando en sus funciones básicas si el servicio de notificaciones está temporalmente indisponible.')]:
 h(title); p(txt)

page('4. Implementación, entregables y recepción')
p('Cronograma referencial de 24 semanas desde la orden de inicio y entrega de insumos institucionales. El calendario definitivo debe aprobarse antes de publicar y responder al número de sedes y volumen real de migración.')
table(['Hito / plazo','Entregable y condición de aceptación','Pago'],[
['D1 · semanas 1–3','Diagnóstico, alcance detallado, matriz RF/RT, riesgos, plan de proyecto y diseño aprobados.','10%'],
['D2 · semanas 4–10','Ambiente de pruebas, configuración, módulos y versión móvil alfa para Android e iOS; pruebas funcionales trazables aprobadas.','20%'],
['D3 · semanas 11–15','Migración de ensayo, interfaces incluidas y conciliación de datos; pruebas técnicas aprobadas.','20%'],
['D4 · semanas 16–19','Piloto web y móvil, formación por roles, pruebas en dispositivos y aceptación de usuarios; incidencias críticas resueltas.','20%'],
['D5 · semanas 20–24','Despliegue, aplicaciones móviles firmadas y distribuidas, estabilización mínima de 30 días, documentación, transferencia y acta final.','30%']], [1900,6500,960])
p('Los porcentajes corresponden al componente de implementación, suman 100% y se devengan contra entregables aceptados. El soporte recurrente, si tiene precio separado, se pagará por períodos vencidos con informe de servicio. No habrá doble cobro por correcciones cubiertas por garantía.')
h('Migración y puesta en operación')
p('Inventariar fuentes, mapear campos, acordar reglas de depuración y ejecutar al menos una migración de ensayo. Conciliar cantidades por entidad, saldos de inventario, relaciones y muestras clínicas aprobadas; documentar registros rechazados y su resolución. No eliminar las fuentes originales antes de la aceptación y la autorización institucional.')
p('El plan de corte incluirá responsables, respaldo, ventana de indisponibilidad, validaciones, criterios de reversión y procedimiento operativo durante la contingencia. La institución designará responsables clínicos, farmacia, estadística y tecnología para validar cada componente.')
h('Procedimiento de aceptación y cambios')
p('La entidad dispondrá de 10 días hábiles propuestos para aceptar u observar cada entregable por escrito. El proveedor subsanará dentro del plazo acordado según severidad; se repetirá la prueba afectada. El silencio no constituye aceptación. Las variaciones de alcance requieren análisis de impacto en costo, tiempo y servicio y autorización escrita conforme al contrato.')
h('Capacitación y documentación')
p('Entregar manuales de usuario y administración, guía de despliegue, diccionario de datos, API, respaldo y recuperación. Para el componente móvil, incluir guía de publicación, matriz de compatibilidad, configuración por ambiente, certificados y llaves bajo custodia acordada, política de versiones, soporte al usuario y material breve para pacientes. Capacitar por rol y formar administradores institucionales; entregar asistencia y ejercicios prácticos. Meta propuesta: al menos 80% de logro en los ejercicios de cada participante, con refuerzo para quienes no alcancen el resultado.')

page('5. Garantía, soporte y condiciones de operación')
p('Se propone garantía de 12 meses a partir de la recepción final para corregir defectos del alcance aceptado, sin costo adicional. El oferente incluirá mesa de ayuda, canales de contacto, escalamiento y responsables, distinguiendo garantía correctiva de nuevas funcionalidades.')
table(['Severidad','Ejemplo','Respuesta / restauración propuesta'],[
['P1 · crítica','Caída general, exposición de datos o imposibilidad de atención esencial.','30 minutos / 4 horas; cobertura 24×7.'],
['P2 · alta','Módulo esencial inutilizable con contingencia limitada.','2 horas / 8 horas; cobertura 24×7.'],
['P3 · media','Falla parcial con alternativa operativa.','8 horas hábiles / 3 días hábiles.'],
['P4 · baja','Defecto menor sin interrupción del servicio.','2 días hábiles / 10 días hábiles.']], [1500,4200,3660])
p('Restauración significa recuperación del servicio o una contingencia aceptada; no equivale a corrección definitiva. Para P1/P2 se entregará análisis causal y plan de solución permanente dentro de 5 días hábiles. Horario hábil propuesto: lunes a viernes, 8:00–17:00, hora de Honduras, excepto feriados.')
h('Medición y seguimiento')
p('La disponibilidad será 100 × (minutos de la ventana contratada menos minutos de indisponibilidad atribuible al servicio) / minutos de la ventana contratada. Solo se excluirán mantenimientos autorizados previamente y causas externas documentadas según el contrato. Cada incidente tendrá hora de apertura, clasificación, respuesta, restauración y cierre validado.')
p('Presentar informe mensual de disponibilidad, incidentes, cumplimiento de tiempos, respaldos y medidas correctivas. Los descuentos, penalidades, límites y garantías contractuales se fijarán en las condiciones administrativas antes de convocar, evitando fórmulas o montos no aprobados.')
h('Licencias, derechos y entrega al finalizar')
p('La entidad conservará control y acceso a sus datos. El oferente identificará componentes propios y de terceros, restricciones, costos recurrentes y derechos de uso. El pliego deberá seleccionar antes de publicarse el régimen de licenciamiento, alojamiento y derechos sobre desarrollos específicos; no se presume transferencia del código preexistente.')
p('Entregar las configuraciones y desarrollos específicos según los derechos contratados. Cuando la continuidad requiera acceso al código preexistente, definir licencia de mantenimiento o custodia con causales y alcance. Al terminar, proporcionar datos, adjuntos, documentación y apoyo de transición durante un plazo propuesto de 30 días, sin bloquear exportaciones.')
h('Obligaciones de la entidad')
p('Proveer contrapartes y decisiones oportunas, catálogos y reglas clínicas aprobadas, fuentes de datos autorizadas, conectividad y accesos bajo su responsabilidad. El proveedor mantendrá una matriz de dependencias con alertas tempranas y propuestas de mitigación.')

page('6. Contenido de oferta y evaluación propuesta')
p('La unidad de contratación deberá confirmar que el método de selección y la ponderación son admisibles para la modalidad y fuente de financiamiento. La siguiente matriz es una propuesta técnica; no reemplaza el método legal aplicable ni el pliego administrativo.')
h('Documentación de la oferta')
p('Presentar resumen de solución; matriz RF/RT/RM con estado y evidencia; arquitectura; plan y equipo de implementación; cronograma; migración; capacitación; garantía y soporte; licencias y dependencias; riesgos y exclusiones; y oferta económica desglosada. La documentación de habilitación, solvencia y garantías será la exigida en las bases del proceso.')
p('Equipo propuesto: coordinación de proyecto, análisis de procesos de salud, desarrollo o configuración, migración e integración, aseguramiento de calidad y capacitación/soporte. Se admitirán roles combinados cuando la dedicación y capacidad estén justificadas. Acreditar experiencia comparable con contratos, actas o referencias verificables, sin exigir experiencia con una entidad o marca específica.')
table(['Componente','Puntos','Regla propuesta'],[
['Prueba funcional','40','Diez escenarios de la sección 7, 4 puntos cada uno.'],
['Solidez técnica','15','Seguridad, restauración y rendimiento: 5 puntos cada uno.'],
['Implementación','10','Plan/migración y capacitación/transferencia: 5 puntos cada uno.'],
['Equipo y experiencia','5','Experiencia comparable verificable: 2.5; equipo y dedicación: 2.5.'],
['Costo total evaluado','30','30 × menor costo admisible / costo de la oferta evaluada.'],
['Total','100','Aplicar únicamente si se aprueba este método.']], [2300,650,6410])
p('Para cada subcriterio técnico de 5 puntos: 5 si acredita todos sus elementos; 2.5 si presenta evidencia parcial y un plan de cierre verificable; 0 si no acredita. Para cada subcriterio de 2.5: 2.5 si acredita todos sus elementos; 1.25 si es parcial; 0 sin evidencia. El protocolo definitivo detallará los elementos exigidos antes de publicar.')
p('Umbral técnico propuesto: 49/70, además de los requisitos de admisibilidad aprobados. Todos los requisitos RF, RT y RM deberán cumplirse al recibir la solución. Si una capacidad debe existir al ofertar, deberá señalarse expresamente y justificarse antes de la convocatoria. La evaluación no podrá agregar requisitos durante la demostración.')
h('Costo comparable')
p('Comparar un horizonte propuesto de 36 meses: implementación, licencias, alojamiento, soporte, integraciones y componentes obligatorios, con iguales volúmenes y tratamiento tributario. Desglosar moneda, impuestos, renovaciones y precios unitarios por crecimiento. El plazo contractual y presupuesto deben ser compatibles con el horizonte adoptado; evitar contar dos veces el soporte incluido en garantía.')

page('7. Demostración y decisiones previas a convocatoria')
p('Aplicar el mismo guion, datos sintéticos, tiempo y ambiente de evaluación a todos los oferentes. Tiempo propuesto: 90 minutos más 30 de preguntas. Registrar resultados en acta; las capturas o presentaciones no sustituyen la ejecución cuando se evalúe una función operativa.')
table(['Caso','Evidencia observable para puntaje completo'],[
['E1 · acceso y red','Usuario autorizado en dos sedes; acceso a una tercera rechazado; registro de auditoría.'],
['E2 · agenda','Crear jornada y excepción; bloquear doble reserva entre sedes; reprogramar con trazabilidad.'],
['E3 · atención','Identificar paciente, registrar triaje y SOAP, diagnóstico, receta y orden en un mismo historial.'],
['E4 · farmacia','Seleccionar lote FEFO; realizar entrega parcial; registrar saldo y demanda insatisfecha; impedir duplicación del descargo.'],
['E5 · vacunación','Registrar dosis y lote, actualizar saldo y generar carnet vinculado al paciente.'],
['E6 · epidemiología','Activar alerta por diagnóstico configurado, completar ficha y consultar consolidado restringido.'],
['E7 · continuidad','Registrar ingreso, traslado y egreso; consultar controles maternos o infantiles y resultado diagnóstico.'],
['E8 · acceso móvil','Solicitar cuenta, iniciar sesión, forzar cambio de contraseña, cerrar sesión y rechazar reutilización del token.'],
['E9 · servicios móviles','Reservar cita; consultar historial y receta; buscar disponibilidad con fecha de actualización; mostrar carnet verificable.'],
['E10 · seguridad y resiliencia móvil','Rechazar acceso a otro paciente, proteger notificación y archivo, interrumpir/reanudar red sin duplicar la operación y demostrar paquete de producción firmado.']], [1800,7560])
h('Datos y acuerdos por completar')
p('Antes de publicar: confirmar entidad y modalidad; financiamiento y presupuesto; sedes y carga; versiones mínimas de Android e iOS; dispositivos objetivo; operación sin conexión; tiendas o distribución institucional; cuentas de publicación; mensajería y costos recurrentes; interfaces; alojamiento; volumen de migración; plazos; matriz de evaluación; licencias y derechos; régimen de soporte; responsables de aceptación y condiciones administrativas. Validar con las áreas clínica, informática, compras y asesoría jurídica.')
h('Referencia para contratación en Honduras')
p('Se ha utilizado Honduras como contexto preliminar por la documentación del SISS. El texto de la Ley de Contratación del Estado publicado por ONCAE, artículo 41, orienta a preparar especificaciones que favorezcan la competencia y a evitar exigencias técnicamente innecesarias que restrinjan la participación. Las bases deben admitir equivalencia funcional y criterios objetivos.')
p('Fuente oficial consultada: ONCAE, Ley de Contratación del Estado, versión revisada diciembre de 2016, artículo 41: https://oncae.gob.hn/wp-content/uploads/2024/04/Ley-de-Contratacion-del-Estado-revisada-diciembre-2016.pdf')
p('Consulta: 7 de septiembre de 2026. La unidad jurídica confirmará reformas, disposiciones presupuestarias, documentos estándar y reglas del financiador aplicables al proceso concreto; esta referencia no determina la modalidad ni los umbrales de contratación.')

page('Anexo interno · preparación competitiva del SISS')
p('USO INTERNO. Retirar este anexo antes de presentar o publicar los términos de referencia. La preparación competitiva consiste en llegar con evidencia reproducible, alcance claro y costos sustentados.')
h('Capacidades que conviene demostrar primero')
p('Priorizar una demostración continua desde la cita hasta el expediente, la receta y el descargo de inventario. Preparar casos de FEFO, entregas parciales y demanda insatisfecha; rotación de profesionales entre establecimientos; vacunación con lote y carnet; alerta epidemiológica; y consulta personal desde el canal ciudadano.')
h('Base documental revisada')
p('README.md y SISS_Documentacion_Sistema.md describen atención SOAP/CIE-10, agendas, farmacia, vacunación, epidemiología, apoyo diagnóstico y canal móvil. DOCUMENTACION_TECNICA.md describe API, autenticación, auditoría y operaciones transaccionales. El código Flutter contiene flujos para solicitud y acceso de pacientes, citas, recetas, historial clínico, búsqueda de existencias, carnet con QR y controles prenatales con descarga de PDF. Estas afirmaciones deben contrastarse con pruebas antes de incorporarlas como capacidades acreditadas en una oferta.')
p('Las rutas siss-frontend/src/app/app.routes.ts y los directorios siss-backend/src/modules incluyen hospitalización, control prenatal, pediatría, formularios, referencias y otros módulos. siss-mobile/lib contiene pantallas de citas, farmacia, carnet, historial, stock y control prenatal. La revisión realizada fue documental y estructural; no incluyó pruebas de ejecución ni auditoría integral del código.')
h('Brechas que deben cerrarse antes de prometer cumplimiento')
p('Verificar funcionamiento extremo a extremo y permisos en API; bloqueo de duplicados y concurrencia; protección de bitácoras; restauración; pruebas de carga; autenticación multifactor; cifrado; versionado de formularios; distribución móvil e integraciones externas. En la app móvil deben sustituirse la URL local y el tráfico HTTP por configuración HTTPS de producción; mover tokens desde preferencias generales a almacenamiento seguro; configurar firma e identificador institucional; implementar y probar notificaciones; revisar permisos de ubicación; proteger PDF y QR; eliminar datos sensibles de registros; y completar pruebas Android e iOS. No afirmar integración productiva con PACS, analizadores o registros nacionales sin una prueba y autorización de la contraparte.')
h('Paquete recomendado para participar')
p('Preparar una matriz RF/RT/RM con evidencia y responsable, una instalación demostrable con datos sintéticos, paquetes móviles firmados de prueba, matriz de dispositivos, actas de pruebas fechadas, fichas del equipo, referencias autorizadas, cronograma con dependencias y costos totales a 36 meses. Para cada desarrollo pendiente, indicar esfuerzo, precio, fecha y criterio de aceptación.')
p('Revisar especialmente los objetivos propuestos de disponibilidad, soporte, concurrencia y 24 semanas: no provienen de mediciones del SISS y pueden requerir inversión adicional. Diferenciar lo disponible de lo que se ofrece desarrollar evita compromisos difíciles de cumplir y permite defender técnicamente el valor de la propuesta.')
h('Control de liberación')
p('Responsable técnico: [COMPLETAR]. Revisión de compras y jurídica: [COMPLETAR]. Fecha de aprobación: [COMPLETAR]. Retirar este anexo y resolver todos los campos pendientes antes de emitir una versión institucional.')

d.core_properties.title='Términos de referencia — Plataforma integral de servicios de salud'
d.core_properties.subject='Borrador técnico de contratación basado en el alcance SISS'
d.core_properties.author=''
d.save(OUT/'TDR_SISS_Licitacion.docx')
print(OUT/'TDR_SISS_Licitacion.docx')
